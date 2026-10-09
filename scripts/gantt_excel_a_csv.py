"""Convierte la carta Gantt en Excel (formato semanas en columnas) al CSV de programa de la app.

Formato del Excel (Gantt Casas Loma La Cruz):
  - una fila con la fecha del lunes de cada semana y, justo debajo, otra con el viernes;
  - filas de tareas con el nombre en la columna B y, en cada semana, los sitios que se trabajan:
    "6-5" = sitios 6 y 5, "25" = sitio 25.

Uso:
    python scripts/gantt_excel_a_csv.py Gantt.xlsx data/carga_inicial/5_programa_gantt.csv

Por cada sitio el programa va del lunes de su primera semana al viernes de su última semana, considerando
todas las tareas del Excel (de "Entrega de plataformas" a "Relleno compactado"). Con --desde-tarea se puede
partir desde otra tarea (ej. --desde-tarea "Excavación"). También escribe <salida>_detalle.csv con cada
tarea y semana, como respaldo.
"""
import argparse
import csv
import re
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl


def a_fecha(v):
    if isinstance(v, datetime):
        return v.date()
    return v if isinstance(v, date) else None


def sitios_de(v):
    if v is None:
        return []
    if isinstance(v, (int, float)):
        return [int(v)]
    return [int(x) for x in re.findall(r'\d+', str(v)) if 1 <= int(x) <= 48]


def leer(xlsx):
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    # fila de lunes: la que tiene más fechas; la de viernes es la siguiente
    conteo = {r: sum(a_fecha(c.value) is not None for c in ws[r]) for r in range(1, min(ws.max_row, 30) + 1)}
    fila_ini = max(conteo, key=conteo.get)
    semanas = {}
    for c in range(1, ws.max_column + 1):
        ini, fin = a_fecha(ws.cell(fila_ini, c).value), a_fecha(ws.cell(fila_ini + 1, c).value)
        if ini and fin:
            semanas[c] = (ini, fin)
    # columnas a la derecha sin fecha en el encabezado: se siguen numerando semana a semana
    ultima = max(semanas)
    for c in range(ultima + 1, ws.max_column + 1):
        ini, fin = semanas[c - 1]
        semanas[c] = (ini + timedelta(days=7), fin + timedelta(days=7))
    tareas = []
    for r in range(fila_ini + 2, ws.max_row + 1):
        nombre = ws.cell(r, 2).value
        if not isinstance(nombre, str) or not nombre.strip():
            continue
        celdas = [(semanas[c], sitios_de(ws.cell(r, c).value)) for c in semanas if sitios_de(ws.cell(r, c).value)]
        if celdas:
            tareas.append((r, nombre.strip(), celdas))
    return tareas


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('salida')
    ap.add_argument('--desde-tarea', help='usar solo las tareas desde la primera cuyo nombre contenga este texto')
    a = ap.parse_args()
    tareas = leer(a.xlsx)
    if a.desde_tarea:
        i = next(i for i, t in enumerate(tareas) if a.desde_tarea.lower() in t[1].lower())
        tareas = tareas[i:]
    detalle, rango = [], {}
    for _, nombre, celdas in tareas:
        for (ini, fin), sitios in celdas:
            for s in sitios:
                detalle.append((s, nombre, ini, fin))
                i0, f0 = rango.get(s, (ini, fin))
                rango[s] = (min(i0, ini), max(f0, fin))
    salida = Path(a.salida)
    with salida.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['sitio', 'inicio', 'termino', 'tarea'])
        for s in sorted(rango, key=lambda s: rango[s]):
            ini, fin = rango[s]
            w.writerow([s, ini.isoformat(), fin.isoformat(), f'{tareas[0][1]} → {tareas[-1][1]}'])
    with salida.with_name(salida.stem + '_detalle.csv').open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['sitio', 'tarea', 'inicio', 'termino'])
        for s, nombre, ini, fin in sorted(detalle, key=lambda x: (x[0], x[2])):
            w.writerow([s, nombre, ini.isoformat(), fin.isoformat()])
    print(f'{len(rango)} sitios, {len(tareas)} tareas ({tareas[0][1]} … {tareas[-1][1]}) → {salida}')


if __name__ == '__main__':
    main()
